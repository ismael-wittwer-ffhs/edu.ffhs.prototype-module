#!/usr/bin/env python3
"""Extract plain text from binary document formats (PDF, DOCX, PPTX, XLSX).

Scans src/input/ for supported files and writes a .txt file alongside each
original. These text files can then be consumed by the transform-sources
Cursor command.

Usage:
    python scripts/extract-text.py [--force] [input_dir]
"""

import argparse
import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx"}


# --- extractors ---------------------------------------------------------- #

def extract_pdf(path: Path) -> str:
    import pymupdf
    parts: list[str] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n".join(parts)


def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            parts.append(f"--- Slide {slide_num} ---\n" + "\n".join(texts))
    return "\n\n".join(parts)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            parts.append(f"--- {sheet.title} ---\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
}


# --- main ----------------------------------------------------------------- #

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract text from binary documents in src/input/.",
    )
    parser.add_argument(
        "input_dir",
        nargs="?",
        default="src/input",
        help="Directory to scan (default: src/input)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Overwrite existing .txt files",
    )
    args = parser.parse_args()

    input_dir = Path(args.input_dir)
    if not input_dir.is_dir():
        print(f"Error: {input_dir} is not a directory.", file=sys.stderr)
        sys.exit(1)

    files = sorted(
        f for f in input_dir.rglob("*") if f.suffix.lower() in SUPPORTED
    )

    if not files:
        print(f"No supported files found in {input_dir}.")
        return

    extracted, skipped = 0, 0
    for filepath in files:
        txt_path = filepath.with_suffix(".txt")
        if txt_path.exists() and not args.force:
            print(f"  skip  {filepath.relative_to(input_dir)}  (txt exists)")
            skipped += 1
            continue

        extractor = EXTRACTORS[filepath.suffix.lower()]
        try:
            text = extractor(filepath)
            txt_path.write_text(text, encoding="utf-8")
            print(f"  done  {filepath.relative_to(input_dir)} -> {txt_path.name}")
            extracted += 1
        except Exception as exc:
            print(f"  FAIL  {filepath.relative_to(input_dir)}: {exc}", file=sys.stderr)

    print(f"\nExtracted: {extracted}  Skipped: {skipped}  Total: {len(files)}")


if __name__ == "__main__":
    main()
